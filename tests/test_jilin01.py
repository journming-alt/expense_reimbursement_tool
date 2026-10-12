from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from workflow import generate_outputs, import_trip_source, scan_trip  # noqa: E402


class Jilin01DemoTest(unittest.TestCase):
    def setUp(self) -> None:
        self.jilin_zip = Path(
            os.environ.get(
                "JILIN_ZIP_PATH",
                "/Volumes/Users/journ/Desktop/win发票/朱艳明05月份差旅报销/吉林01.zip",
            )
        )
        if not self.jilin_zip.exists():
            self.skipTest(f"吉林01.zip 不存在：{self.jilin_zip}")
        work_parent = PROJECT_ROOT / "data" / "test_work"
        work_parent.mkdir(parents=True, exist_ok=True)
        self.work_dir = Path(tempfile.mkdtemp(prefix="jilin01_", dir=work_parent))

    def test_jilin01_end_to_end(self) -> None:
        trip_root = import_trip_source(self.jilin_zip, self.work_dir)
        self.assertEqual(trip_root.name, "吉林01")

        trip = scan_trip(trip_root, examiner="朱艳明", site="吉林01")
        by_number = {invoice.invoice_number: invoice for invoice in trip.invoices}

        train = by_number["26219153885001205196"]
        self.assertEqual(train.passenger_name, "朱艳明")
        self.assertEqual(train.departure_date, "2026-05-07")
        self.assertIn(train.departure_place, {"沈阳", "沈阳北"})
        self.assertIn(train.arrival_place, {"长春", "长春西"})
        self.assertEqual(train.total_amount, 148.0)

        flight = by_number["26118999111035480353"]
        self.assertEqual(flight.passenger_name, "朱艳明")
        self.assertTrue("长春" in (flight.departure_place or "") or "龙嘉" in (flight.departure_place or ""))
        self.assertTrue("成都" in (flight.arrival_place or "") or "天府" in (flight.arrival_place or ""))
        self.assertEqual(flight.vehicle_no, "CA2722")
        self.assertEqual(flight.departure_date, "2026-05-09")
        self.assertEqual(flight.total_amount, 1100.0)

        hotel = by_number["26222000000401165971"]
        self.assertEqual(hotel.hotel_guest, "朱艳明")
        self.assertEqual(hotel.checkin_date, "2026-05-07")
        self.assertEqual(hotel.checkout_date, "2026-05-09")
        self.assertEqual(hotel.nights, 2)
        self.assertEqual(hotel.total_amount, 660.0)

        self.assertTrue(all(invoice.verification_status_label == "已查验" for invoice in trip.invoices))

        output_dir = PROJECT_ROOT / "data" / "test_output_jilin01"
        trip = generate_outputs(trip, output_dir)

        self.assertTrue(Path(trip.generated_excel_path).exists())
        self.assertTrue(Path(trip.exported_zip_path).exists())

        ws = load_workbook(trip.generated_excel_path, data_only=False).active
        self.assertEqual(ws["B4"].value, "考试员：朱艳明")
        self.assertEqual(ws["B5"].value, "考点：吉林01")
        self.assertEqual(ws["H9"].value, 148)
        self.assertEqual(ws["H10"].value, 1100)
        self.assertEqual(ws["O9"].value, 660)
        self.assertEqual(ws["H16"].value, "=SUM(H9:I15)")
        self.assertEqual(ws["O12"].value, "=SUM(O9:O11)")
        self.assertEqual(ws["C22"].value, "=H16+O12+H21+N16")

        self.assertEqual(sum(ws.cell(row, 8).value or 0 for row in range(9, 16)), 1248)
        self.assertEqual(ws["O9"].value, 660)
        self.assertEqual(1248 + 660, 1908)


if __name__ == "__main__":
    unittest.main(verbosity=2)
