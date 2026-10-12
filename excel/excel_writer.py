from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from excel.template_detector import TemplateLayout, detect_template_from_sheet
from models.invoice import Invoice
from models.trip import TripContext
from utils.date_utils import display_date
from utils.money_utils import money_or_zero, round_money


class ExcelWriter:
    def __init__(self, template_path: str | Path):
        self.template_path = Path(template_path)

    def write(self, trip: TripContext, output_path: str | Path) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.template_path, output_path)

        wb = load_workbook(output_path)
        ws = wb.active
        layout = detect_template_from_sheet(ws)

        reimbursable = [i for i in trip.invoices if i.document_type != "support_doc"]
        transports = sorted(
            [i for i in reimbursable if i.category == "transport"],
            key=lambda i: (i.departure_date or "", i.file_name),
        )
        hotels = sorted([i for i in reimbursable if i.category == "hotel"], key=lambda i: (i.checkin_date or "", i.file_name))
        meals = [i for i in reimbursable if i.category == "meal"]
        local_transports = [i for i in reimbursable if i.category == "local_transport"]

        layout = self._ensure_rows(ws, layout, "hotel", len(hotels))
        layout = self._ensure_rows(ws, layout, "transport", len(transports))
        layout = self._ensure_rows(ws, layout, "meal", len(meals))
        layout = self._ensure_rows(ws, layout, "local", len(local_transports))

        self._write_header(ws, layout, trip)
        self._clear_data(ws, layout)
        self._write_transports(ws, layout, transports, trip.examiner)
        self._write_hotels(ws, layout, hotels)
        self._write_meals(ws, layout, meals)
        self._write_local_transports(ws, layout, local_transports)
        self._write_formulas(ws, layout)

        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.calculation.calcMode = "auto"
        except Exception:
            pass

        wb.save(output_path)
        return output_path

    def _write_header(self, ws: Worksheet, layout: TemplateLayout, trip: TripContext) -> None:
        self._set(ws, layout.examiner_cell, f"考试员：{trip.examiner}")
        self._set(ws, layout.site_cell, f"考点：{trip.site}")

    def _clear_data(self, ws: Worksheet, layout: TemplateLayout) -> None:
        for row in range(layout.transport_start_row, layout.transport_subtotal_row):
            for col in range(2, 9):
                self._clear_cell(ws, row, col)
        for row in range(layout.hotel_start_row, layout.hotel_subtotal_row):
            for col in (10, 12, 13, 14, 15):
                self._clear_cell(ws, row, col)
        for row in range(layout.meal_start_row, layout.meal_subtotal_row):
            for col in (10, 14):
                self._clear_cell(ws, row, col)
        for row in range(layout.local_start_row, layout.local_subtotal_row):
            for col in (3, 10):
                self._clear_cell(ws, row, col)

    def _write_transports(self, ws: Worksheet, layout: TemplateLayout, invoices: list[Invoice], examiner: str) -> None:
        for index, invoice in enumerate(invoices):
            row = layout.transport_start_row + index
            name = invoice.passenger_name or examiner
            self._set(ws, row, 2, name)
            self._set(ws, row, 3, display_date(invoice.departure_date, full=index == 0))
            self._set(ws, row, 4, _short_place(invoice.departure_place))
            self._set(ws, row, 5, display_date(invoice.arrival_date or invoice.departure_date, full=index == 0))
            self._set(ws, row, 6, _short_place(invoice.arrival_place))
            self._set(ws, row, 7, invoice.transport_tool or "")
            self._set(ws, row, 8, money_or_zero(invoice.total_amount))

    def _write_hotels(self, ws: Worksheet, layout: TemplateLayout, invoices: list[Invoice]) -> None:
        for index, invoice in enumerate(invoices):
            row = layout.hotel_start_row + index
            self._set(ws, row, 10, _short_place(invoice.hotel_city))
            self._set(ws, row, 12, money_or_zero(invoice.unit_price))
            self._set(ws, row, 13, invoice.rooms or 1)
            self._set(ws, row, 14, invoice.nights or "")
            self._set(ws, row, 15, money_or_zero(invoice.total_amount))

    def _write_meals(self, ws: Worksheet, layout: TemplateLayout, invoices: list[Invoice]) -> None:
        for index, invoice in enumerate(invoices):
            row = layout.meal_start_row + index
            self._set(ws, row, 10, invoice.seller_name or invoice.invoice_type)
            self._set(ws, row, 14, money_or_zero(invoice.total_amount))

    def _write_local_transports(self, ws: Worksheet, layout: TemplateLayout, invoices: list[Invoice]) -> None:
        for index, invoice in enumerate(invoices):
            row = layout.local_start_row + index
            desc = _local_transport_desc(invoice)
            self._set(ws, row, 3, desc)
            self._set(ws, row, 10, money_or_zero(invoice.total_amount))

    def _write_formulas(self, ws: Worksheet, layout: TemplateLayout) -> None:
        self._set(ws, layout.transport_subtotal_row, 8, f"=SUM(H{layout.transport_start_row}:I{layout.transport_subtotal_row - 1})")
        self._set(ws, layout.hotel_subtotal_row, 15, f"=SUM(O{layout.hotel_start_row}:O{layout.hotel_subtotal_row - 1})")
        self._set(ws, layout.meal_subtotal_row, 14, f"=SUM(N{layout.meal_start_row}:O{layout.meal_subtotal_row - 1})")
        self._set(ws, layout.local_subtotal_row, 8, f"=SUM(J{layout.local_start_row}:O{layout.local_subtotal_row - 1})")
        self._set(
            ws,
            layout.total_row,
            3,
            f"=H{layout.transport_subtotal_row}+O{layout.hotel_subtotal_row}+H{layout.local_subtotal_row}+N{layout.meal_subtotal_row}",
        )

    def _ensure_rows(self, ws: Worksheet, layout: TemplateLayout, section: str, needed: int) -> TemplateLayout:
        if needed <= 0:
            return layout
        if section == "hotel":
            start, subtotal = layout.hotel_start_row, layout.hotel_subtotal_row
        elif section == "transport":
            start, subtotal = layout.transport_start_row, layout.transport_subtotal_row
        elif section == "meal":
            start, subtotal = layout.meal_start_row, layout.meal_subtotal_row
        else:
            start, subtotal = layout.local_start_row, layout.local_subtotal_row

        capacity = max(0, subtotal - start)
        if needed <= capacity:
            return layout

        extra = needed - capacity
        insert_at = subtotal
        source_row = max(start, subtotal - 1)
        ws.insert_rows(insert_at, extra)
        for offset in range(extra):
            target_row = insert_at + offset
            self._copy_row_style(ws, source_row, target_row)
            self._copy_row_merges(ws, source_row, target_row)
        return detect_template_from_sheet(ws)

    def _copy_row_style(self, ws: Worksheet, source_row: int, target_row: int) -> None:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
        for col in range(1, ws.max_column + 1):
            source = ws.cell(source_row, col)
            target = ws.cell(target_row, col)
            if isinstance(target, MergedCell):
                continue
            if source.has_style:
                target._style = copy(source._style)
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
            target.protection = copy(source.protection)

    def _copy_row_merges(self, ws: Worksheet, source_row: int, target_row: int) -> None:
        existing = {str(rng) for rng in ws.merged_cells.ranges}
        for merged in list(ws.merged_cells.ranges):
            if merged.min_row == source_row and merged.max_row == source_row:
                start = f"{get_column_letter(merged.min_col)}{target_row}"
                end = f"{get_column_letter(merged.max_col)}{target_row}"
                ref = f"{start}:{end}"
                if ref not in existing:
                    try:
                        ws.merge_cells(ref)
                    except ValueError:
                        pass

    def _set(self, ws: Worksheet, cell_or_row: str | int, col: int | object | None = None, value: object | None = None) -> None:
        if isinstance(cell_or_row, str):
            cell = ws[cell_or_row]
            new_value = col
        else:
            cell = ws.cell(row=cell_or_row, column=int(col))
            new_value = value
        if isinstance(cell, MergedCell):
            cell = self._top_left_for_merged_cell(ws, cell.row, cell.column)
        cell.value = new_value
        if isinstance(new_value, (int, float)) and not isinstance(new_value, bool):
            cell.number_format = "0.00" if isinstance(new_value, float) and round_money(new_value) != int(new_value) else "0"

    def _clear_cell(self, ws: Worksheet, row: int, col: int) -> None:
        cell = ws.cell(row=row, column=col)
        if not isinstance(cell, MergedCell):
            cell.value = None

    def _top_left_for_merged_cell(self, ws: Worksheet, row: int, col: int):
        for merged in ws.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                return ws.cell(merged.min_row, merged.min_col)
        return ws.cell(row, col)


def _short_place(value: str | None) -> str:
    if not value:
        return ""
    text = str(value).strip()
    if " " in text:
        text = text.split()[0]
    text = text.replace(" ", "")
    for suffix in ("站", "机场"):
        if text.endswith(suffix):
            text = text[: -len(suffix)]
    if text.endswith("市"):
        text = text[:-1]
    return text


def _local_transport_desc(invoice: Invoice) -> str:
    if invoice.departure_place and invoice.arrival_place:
        return f"{invoice.departure_place} -> {invoice.arrival_place}"
    start = invoice.extra.get("taxi_trip_start")
    end = invoice.extra.get("taxi_trip_end")
    if start and end:
        return f"{start} -> {end}"
    return invoice.seller_name or invoice.invoice_type or invoice.file_name
