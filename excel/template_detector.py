from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class TemplateLayout:
    examiner_cell: str = "B4"
    site_cell: str = "B5"
    transport_start_row: int = 9
    transport_subtotal_row: int = 16
    hotel_start_row: int = 9
    hotel_subtotal_row: int = 12
    meal_header_row: int = 13
    meal_start_row: int = 14
    meal_subtotal_row: int = 16
    local_header_row: int = 17
    local_start_row: int = 18
    local_subtotal_row: int = 21
    total_row: int = 22


def detect_template(path: str | Path) -> TemplateLayout:
    wb = load_workbook(path, data_only=False)
    return detect_template_from_sheet(wb.active)


def detect_template_from_sheet(ws: Worksheet) -> TemplateLayout:
    layout = TemplateLayout()

    for row in range(1, ws.max_row + 1):
        b_value = _cell_text(ws, row, 2)
        j_value = _cell_text(ws, row, 10)

        if b_value.startswith("考试员"):
            layout.examiner_cell = f"B{row}"
        elif b_value.startswith("考点"):
            layout.site_cell = f"B{row}"
        elif b_value == "姓名":
            layout.transport_start_row = row + 1
            layout.hotel_start_row = row + 1
        elif "市内交通费" in b_value:
            layout.local_header_row = row
            layout.local_start_row = row + 1
        elif b_value == "合计":
            layout.total_row = row

        if j_value == "餐费":
            layout.meal_header_row = row
            layout.meal_start_row = row + 1

    b_subtotals = []
    j_subtotals = []
    for row in range(1, ws.max_row + 1):
        b_value = _cell_text(ws, row, 2)
        j_value = _cell_text(ws, row, 10)
        if "小" in b_value and "计" in b_value and row > layout.transport_start_row:
            b_subtotals.append(row)
        if j_value == "小计" and row > layout.hotel_start_row:
            j_subtotals.append(row)

    if b_subtotals:
        layout.transport_subtotal_row = b_subtotals[0]
        after_local = [row for row in b_subtotals if row > layout.local_start_row]
        if after_local:
            layout.local_subtotal_row = after_local[0]

    hotel_subtotals = [row for row in j_subtotals if row < layout.meal_header_row]
    meal_subtotals = [row for row in j_subtotals if row > layout.meal_start_row]
    if hotel_subtotals:
        layout.hotel_subtotal_row = hotel_subtotals[0]
    if meal_subtotals:
        layout.meal_subtotal_row = meal_subtotals[0]

    return layout


def _cell_text(ws: Worksheet, row: int, col: int) -> str:
    value = ws.cell(row=row, column=col).value
    return str(value).strip() if value is not None else ""
